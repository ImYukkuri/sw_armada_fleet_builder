# -*- coding: utf-8 -*-
"""从 Armada_规则简述和任务卡翻译_BY_Remo.xlsx 提取任务卡数据。

产出：
- assets/objectives/text.json   任务卡全部字段（中英对照，重建可再生成）
- assets/objectives/images/*.png 卡图（按任务卡 id 命名）

数据模型（与 data/*.json 同构，build_data.py 会汇总到 data/objectives.json）：
{
  "id": "advanced_gunnery",
  "type": "objective",
  "category": "assault",            # assault | defense | navigation
  "name": "Advanced Gunnery",       # 官方英文卡名
  "source": "Core Set",             # 出自（核心盒 / The Corellian Conflict / ...）
  "setup": "...", "setupEn": "...",
  "special": "...", "specialEn": "...",
  "endOfRound": "...", "endOfRoundEn": "...",
  "endOfGame": "...", "endOfGameEn": "...",
  "score": 15,                      # 胜利得分（'--' 记 null）
  "faq": "...", "faqEn": "...",     # 中英各一列（[FAQ] 前缀已剥离）
  "image": "assets/objectives/images/advanced_gunnery.png"
}

用法：
  python extract_objectives.py
"""
import json
import os
import re
import zipfile

HERE = os.path.dirname(os.path.abspath(__file__))
ASSETS = os.path.join(HERE, "assets")
OBJ_ASSETS = os.path.join(ASSETS, "objectives")
IMG_DIR = os.path.join(OBJ_ASSETS, "images")
XLSX = os.path.join(ASSETS, "Armada_规则简述和任务卡翻译_BY_Remo.xlsx")

# ---- 常量 ----
XDR_NS = "http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing"
A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
CJK = re.compile(r"[一-鿿]")

CATEGORIES = {"assault": "进攻", "defense": "防御", "navigation": "导航"}


def slug(name):
    s = name.strip().lower()
    s = re.sub(r"[^a-z0-9]+", "_", s)
    return s.strip("_")


# 原汉化者漏译的英文游戏术语 → 中文（障碍/生物/关键词）。
# 升级/称号卡名（如 Gunnery Team、Impetuous）是专有名词，保留英文原名。
# 仅对中文字段生效；英文原文字段（*En）保持原样。正则用 (?<![A-Za-z])
# 而非 \b 判界，因为汉字也是 \w，\b 无法在英文与汉字之间成界。
TERM_FIXES = [
    (r"stations?", "空间站"),
    (r"dust fields?", "尘埃云"),
    (r"asteroid fields?", "小行星带"),
    (r"asteroid", "小行星带"),      # 单独出现（如「asteroid或碎片带」）
    (r"debris fields?", "碎片带"),
    (r"gravity rifts?", "引力裂缝"),
    (r"purrgils?", "鲸兽"),
    (r"exogorths?", "太空蛞蝓"),
    (r"strategic", ""),  # 与相邻「策略特性」同义重复，去掉冗余英文
]


def fix_terms(text):
    """汉化中文规则文本里残留的英文游戏术语。"""
    for pat, repl in TERM_FIXES:
        text = re.sub(r"(?<![A-Za-z])" + pat + r"(?![A-Za-z])", repl, text, flags=re.I)
    return text


def split_zh_en(text):
    """按行拆中英：含中文的行归 zh，其余归 en。

    不再以「首个中文行」截断——那会把中文行之后夹带的英文行（与原中文同义
    的重复句）漏进 zh 字段，导致 special 等词条出现中英同义并存。"""
    if not text:
        return "", ""
    lines = [l.strip() for l in text.split("\n") if l.strip()]
    zh = [l for l in lines if CJK.search(l)]
    en = [l for l in lines if not CJK.search(l)]
    return "\n".join(zh), "\n".join(en)


def parse_source(name):
    """名称列形如 'Blockade Run\n（The Corellian Conflict）' → (卡名, 出处)。"""
    lines = [l.strip() for l in name.split("\n") if l.strip()]
    card = lines[0]
    src = "Core Set"
    for l in lines[1:]:
        m = re.search(r"[（(]([^）)]+)[）)]", l)
        if m:
            src = m.group(1).strip()
    return card, src


def extract_images():
    """读取 drawing2.xml 的锚点 -> rId -> 图片文件，按行号（0 基 = Excel 第 row+1 行）返回。"""
    z = zipfile.ZipFile(XLSX)
    import xml.etree.ElementTree as ET
    ns = {"xdr": XDR_NS, "a": A_NS, "r": R_NS}

    # rId -> 图片文件
    rel_xml = z.read("xl/drawings/_rels/drawing2.xml.rels").decode("utf-8")
    rel_root = ET.fromstring(rel_xml)
    rid_map = {}
    for rel in rel_root:
        rid = rel.get("Id")
        target = rel.get("Target")
        if target:
            rid_map[rid] = os.path.normpath(os.path.join("xl/drawings", target)).replace("\\", "/")

    draw_xml = z.read("xl/drawings/drawing2.xml").decode("utf-8")
    root = ET.fromstring(draw_xml)
    result = {}  # row_index -> 文件路径
    for anchor in root:
        if not anchor.tag.endswith("Anchor"):
            continue
        fr = anchor.find("xdr:from", ns)
        if fr is None:
            continue
        col = int(fr.findtext("xdr:col", "0", ns) or "0")
        row = int(fr.findtext("xdr:row", "0", ns) or "0")
        blip = anchor.find(".//a:blip", ns)
        if blip is None:
            continue
        rid = blip.get("{%s}embed" % R_NS)
        if rid in rid_map:
            result[row] = rid_map[rid]
    return result, z


def main():
    import openpyxl
    wb = openpyxl.load_workbook(XLSX, data_only=True)
    ws = wb.worksheets[1]  # Objective Cards任务卡翻译

    row_imgs, z = extract_images()
    os.makedirs(IMG_DIR, exist_ok=True)

    objectives = []
    for r in range(2, ws.max_row + 1):
        name = ws.cell(r, 2).value
        cat = ws.cell(r, 3).value
        if not name or not cat:
            continue
        card, src = parse_source(str(name))
        cat = cat.strip().lower()

        def field(col):
            # Excel 用 "--" 表示空栏（英文原文空时显示为 "--"），视为空字符串
            v = ws.cell(r, col).value
            return "" if v is None or str(v).strip() == "--" else str(v)

        zh, en = split_zh_en(field(4))   # D SETUP
        setup, setup_en = fix_terms(zh), en
        zh, en = split_zh_en(field(5))   # E SPECIAL RULE
        special, special_en = fix_terms(zh), en
        zh, en = split_zh_en(field(6))   # F END OF ROUND
        end_round, end_round_en = fix_terms(zh), en
        zh, en = split_zh_en(field(7))   # G END OF GAME
        end_game, end_game_en = fix_terms(zh), en

        faq_zh, faq_en = split_zh_en(field(9))   # I FAQ（中英逐段对照）
        faq = fix_terms(faq_zh)
        # 英文 FAQ 行带 "[FAQ]" 前缀标记，去掉（title 由界面静态输出）
        faq_en = re.sub(r"^\[FAQ\]\s*", "", faq_en)

        score_v = ws.cell(r, 8).value   # H SCORE
        score = None
        if isinstance(score_v, (int, float)):
            score = int(score_v)

        oid = slug(card)
        img_rel = row_imgs.get(r - 1)
        img_name = None
        if img_rel:
            ext = os.path.splitext(img_rel)[1] or ".png"
            img_name = oid + ext
            with z.open(img_rel) as fsrc, open(os.path.join(IMG_DIR, img_name), "wb") as fdst:
                fdst.write(fsrc.read())

        objectives.append({
            "id": oid,
            "type": "objective",
            "category": cat,
            "name": card,
            "source": src,
            "setup": setup, "setupEn": setup_en,
            "special": special, "specialEn": special_en,
            "endOfRound": end_round, "endOfRoundEn": end_round_en,
            "endOfGame": end_game, "endOfGameEn": end_game_en,
            "score": score,
            "faq": faq, "faqEn": faq_en,
            "image": "assets/objectives/images/" + img_name if img_name else "",
        })

    out = os.path.join(OBJ_ASSETS, "text.json")
    with open(out, "w", encoding="utf-8") as f:
        json.dump(objectives, f, ensure_ascii=False, indent=1)

    missing_img = [o["id"] for o in objectives if not o["image"]]
    print("任务卡 %d 张 -> %s" % (len(objectives), out))
    print("按分类: assault=%d defense=%d navigation=%d" % (
        sum(1 for o in objectives if o["category"] == "assault"),
        sum(1 for o in objectives if o["category"] == "defense"),
        sum(1 for o in objectives if o["category"] == "navigation")))
    print("无卡图: %s" % (missing_img or "无"))


if __name__ == "__main__":
    main()
